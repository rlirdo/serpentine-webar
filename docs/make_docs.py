# -*- coding: utf-8 -*-
"""產生《教師指引_課程與評量》.docx 與同名 .pdf。

- 題庫、面向、關卡文字一律從 js/content.js 讀取（本機伺服器＋headless Chrome 以 ES module 匯入後轉 JSON），
  確保文件與遊戲內容一致。
- 字型：中文標楷體、英數 Times New Roman；頁尾「NDHU自資系游豐兆製作」＋頁碼。
- PDF 以 Word COM（DispatchEx("Word.Application")，SaveAs2 FileFormat=17）轉出。

用法（repo 根目錄）：PYTHONUTF8=1 python docs/make_docs.py [輸出資料夾]
預設輸出到 docs/。需要 python-docx、pywin32、websocket-client 與本機 Word。
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

VERSION = "1.0.0"
DATE = "2026/10/01"
ONLINE = "https://rlirdo.github.io/serpentine-webar/"
REPO = "https://github.com/rlirdo/serpentine-webar"
CREDIT = "NDHU自資系游豐兆製作"
NAME = "教師指引_課程與評量"
# 與 teacher.html 的課程對照表一致：各關主要面向
LEVEL_FACET = ["知識", "知識", "態度", "知識", "知識", "行動", "態度", "探究"]
# 各關互動（依規格書第 4 節）
LEVEL_PLAY = [
    "連點石塊 3 階段：切片→1–3 cm 碎塊→粗粉",
    "按住〔轉動球磨罐〕蓄力，粒子縮到約 2–4 微米，表面出現發亮缺口",
    "五個選項選對三個：純水、T 微胞粉、氯化鐵（Fe³⁺）",
    "〔投入〕後連點〔攪拌〕；T 微胞使粉體分散，Fe³⁺ 貼到表面",
    "〔滴一滴深層海水〕，點 Mg²⁺ 送進 6 個晶格空位；點到 Ca²⁺ 會提示",
    "〔離心〕分層，再〔洗滌〕×3；選擇廢液處理方式",
    "三選一：50 °C 慢慢烘 48 小時／200 °C 快速烘／直接曬到明年",
    "〔FTIR 檢查〕簡化光譜；比較「修復前／修復後（預期）」4–14 μm 波紋",
]


def load_content():
    import devserver
    from chrome import Chrome
    srv = devserver.start(8771)
    ch = Chrome(port=9351, width=800, height=600)
    try:
        ch.send("Page.enable")
        ch.send("Page.navigate", {"url": "http://127.0.0.1:8771/index.html"})
        import time
        time.sleep(1.0)
        data = ch.eval("""import('/js/content.js').then(m => JSON.stringify({
            LEVELS: m.LEVELS, BANK: m.BANK, FACETS: m.FACETS, FACET_NAMES: m.FACET_NAMES, ENDING: m.ENDING,
            GRADE_LABEL: m.GRADE_LABEL }))""")
    finally:
        ch.close()
        srv.shutdown()
    return json.loads(data)


# ---------------------------------------------------------------- docx 輔助
def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), "標楷體")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def para(doc_or_cell, text="", size=10.5, bold=False, align=None, space_after=2, color=None, style=None):
    p = doc_or_cell.add_paragraph(style=style) if style else doc_or_cell.add_paragraph()
    if text:
        r = p.add_run(text)
        set_run_font(r, size, bold, color)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.15
    if align:
        p.alignment = align
    return p


def heading(doc, text, size=13):
    p = para(doc, text, size=size, bold=True, space_after=3, color="0B1F3A")
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text, size=10.5, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("• ")
    set_run_font(r, size)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        set_run_font(r, size, True)
    r = p.add_run(text)
    set_run_font(r, size)
    return p


def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def table(doc, header, rows, widths, size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, size, True, "FFFFFF")
        shade(c, "065A82")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(v))
            set_run_font(r, size)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    # 標題列跨頁重複
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def add_page_field(paragraph):
    r = paragraph.add_run()
    set_run_font(r, 9)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._element.append(el)


# ---------------------------------------------------------------- 內容
def build(C, out_docx):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run(f"{CREDIT}　｜　第 ")
    set_run_font(r, 9)
    add_page_field(fp)
    r = fp.add_run(" 頁")
    set_run_font(r, 9)

    FN = C["FACET_NAMES"]
    BANK = C["BANK"]
    LEVELS = C["LEVELS"]

    para(doc, "墨玉晶格修復任務 WebAR", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, color="0B1F3A")
    para(doc, "教師指引：課程與評量", size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, color="065A82")
    para(doc, f"版本 {VERSION}｜{DATE}｜線上網址 {ONLINE}", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, color="64748B")

    heading(doc, "一、遊戲簡介與適用對象")
    para(doc, "本遊戲把蛇紋石（墨玉）永續材料研究的「材料端實驗流程八步驟」轉譯成手機 WebAR 闖關：學生用手機掃描任務卡，"
              "主角「墨玉小晶」浮現在卡片上方，帶領學生從切割粉碎、球磨微米化、配製反應液，一路做到深層海水滴定與表徵分析。"
              "原料為在地撿石人無償提供的閒置蛇紋石（墨玉），核心概念是「沒人要的石頭能不能變成有用的材料」。")
    bullet(doc, "國小高年級（五、六年級）選〔國小版〕：用「補破洞」「看不見的光」等比喻。", bold_prefix="適用對象：")
    bullet(doc, "國中（七至九年級）選〔國中版〕：使用表面積、離子交換、晶格、觸媒、波長等用詞。")
    bullet(doc, "一節課 40–45 分鐘；遊戲本身約 8–10 分鐘，其餘為引導、討論與評量。", bold_prefix="時間：")
    bullet(doc, "學生手機或平板（能上網、有相機），以投影或紙本任務卡為掃描目標；沒有手機時可用電腦〔直接玩〕（免相機、免網路）。",
           bold_prefix="設備：")

    heading(doc, "二、學習目標（環境教育素養四面向）")
    facet_goal = {
        "知識": "知道顆粒變小會使表面積變大；認識深層海水是花蓮在地海洋資源；知道遠紅外線是看不見的光（熱輻射）；國中生能說明離子交換與觸媒的角色。",
        "態度": "願意把閒置石料視為可再利用的資源（邊緣資材增值）；認同綠色化學以水為溶劑、低溫省能的做法。",
        "行動": "知道實驗廢液要收集處理、不可倒入水槽；能選擇一項日常可做到的環境行動並做出承諾。",
        "探究": "能區分「科學家的猜想（工作假設）」與「已被證據支持的結論」，知道要用 FTIR／XRD／XPS 與放射率量測取得證據。",
    }
    rows = []
    for f in C["FACETS"]:
        lv = "、".join(str(i + 1) for i, x in enumerate(LEVEL_FACET) if x == f)
        qe = "、".join(q["id"] for q in BANK["elem"] if q["facet"] == f)
        qj = "、".join(q["id"] for q in BANK["junior"] if q["facet"] == f)
        extra = "；行動承諾" if f == "行動" else ("；第 8 關星數" if f == "探究" else "")
        rows.append([FN[f], facet_goal[f], f"第 {lv} 關", f"國小 {qe}\n國中 {qj}{extra}"])
    table(doc, ["面向", "學習目標", "主要關卡", "對應題號"], rows, [2.0, 9.6, 2.0, 3.4])

    heading(doc, "三、一節課流程建議（45 分鐘；40 分鐘版可縮短引起動機與統整）")
    flow = [
        ["0–5", "引起動機", "投影啟動頁（exe）或任務卡，展示蛇紋石照片，提問：「沒人要的石頭，能不能變成有用的材料？」"],
        ["5–10", "準備與前測", "學生掃 QR 開網頁→允許相機→對準任務卡；選國小版／國中版；遊戲自動進行前測 3 題。"],
        ["10–28", "八關闖關", "學生自行闖關。建議在第 5 關（深層海水滴定）與第 8 關（表徵分析）後暫停 1–2 分鐘，全班討論「猜想與證據」。"],
        ["28–33", "後測與承諾", "遊戲自動進行後測 5 題（含前測同 3 題）與行動承諾單選，出現結果頁與證書。"],
        ["33–41", "統整討論", "請學生比較前後測答對數、看四面向長條；教師說明科學誠實聲明（見第八節），分享行動承諾。"],
        ["41–45", "收集評量", "學生在結果頁截圖或舉手回報；共用裝置則由教師開教師頁匯出 CSV（見第六節）。"],
    ]
    table(doc, ["分鐘", "活動", "教師引導重點"], flow, [1.5, 2.3, 13.2])

    heading(doc, "四、八關與實驗流程、學習重點對照表")
    rows = []
    for i, L in enumerate(LEVELS):
        learn = L["learn"]
        rows.append([str(i + 1), L["step"], LEVEL_PLAY[i], f"國小：{learn['elem']}\n國中：{learn['junior']}", FN[LEVEL_FACET[i]]])
    table(doc, ["關", "實驗流程步驟", "遊戲互動", "學習重點（遊戲內文字）", "面向"], rows, [0.8, 2.8, 4.6, 7.0, 1.8], size=8.5)
    para(doc, "每關依操作錯誤次數給 1–3 顆星（錯一次少一顆，最少 1 顆），並附一句「環境小語」。第 5、8 關會明確告訴學生："
              "「這是科學家的猜想，要靠實驗證明」。", size=9.5)

    heading(doc, "五、前後測題目與面向對照")
    para(doc, "前測 3 題（標「前」者）；後測 5 題＝前測同 3 題＋從「行動」面向抽 1 題＋從其他面向抽 1 題（每次隨機）。"
              "選項顯示時會洗牌；下表「正解」為正確選項文字。", size=9.5)
    for g in ("elem", "junior"):
        rows = []
        for q in BANK[g]:
            rows.append([q["id"], "前／後" if q.get("pre") else "後（抽）", FN[q["facet"]], q["q"], q["options"][q["answer"]]])
        para(doc, f"{C['GRADE_LABEL'][g]}題庫", size=10.5, bold=True, space_after=1)
        table(doc, ["題號", "使用", "面向", "題目", "正解"], rows, [1.1, 1.6, 1.9, 8.4, 4.0], size=8.5)

    heading(doc, "六、評量方式")
    bullet(doc, "結果頁顯示「前測 x／3 → 後測同題 y／3」，比較同題進步；另有後測全部 5 題答對數。", bold_prefix="前後測：")
    bullet(doc, "八關共 24 顆星，反映操作中的錯誤嘗試次數，可作為參與度與理解程度的參考。", bold_prefix="星數：")
    bullet(doc, "以後測各面向答對率計算；「行動意向」再與是否完成行動承諾平均，「科學探究」再與第 8 關星數平均。",
           bold_prefix="四面向百分比：")
    bullet(doc, "四選一（" + "、".join(C["ENDING"]["commits"]) + "），教師頁可看全班分布。", bold_prefix="行動承諾：")
    bullet(doc, "在作答的那台裝置開〔教師頁〕（首頁下方連結、結果頁〔給老師看〕或 exe 啟動頁〔教師頁〕）→ 按〔匯出 CSV〕，"
                "下載「墨玉晶格修復任務_作答紀錄_日期.csv」（UTF-8，Excel 可直接開啟）。欄位：序號、時間、版本、前測(3)、後測同題(3)、"
                "後測(5)、星數(24)、四面向%、行動承諾、前後測逐題作答（題號:1 答對／0 答錯）、各關星數。下一班使用前可按〔清除本機紀錄〕。",
           bold_prefix="匯出 CSV：")
    bullet(doc, "紀錄只存在「作答的那台裝置、那個瀏覽器」。學生用自己手機時，請學生於結果頁截圖或口頭回報；"
                "使用共用平板或電腦教室時，每台各自匯出再合併。exe 版的紀錄存在 http://127.0.0.1:8731 之下，"
                "請固定同一台電腦、同一個瀏覽器使用才會累積。", bold_prefix="注意：")
    bullet(doc, "題數少，結果適合作為形成性評量與教學回饋，不宜作為成績的唯一依據。", bold_prefix="解讀：")

    heading(doc, "七、隱私說明")
    bullet(doc, "遊戲不詢問姓名、學號，不需註冊或登入。")
    bullet(doc, "作答結果只存在該裝置瀏覽器的 localStorage，不會上傳；沒有分析追蹤碼，也不連任何外部 API。")
    bullet(doc, "相機畫面只在裝置上即時辨識任務卡，不錄影、不上傳。線上網址為 GitHub Pages 提供的靜態網頁；exe 電腦版完全離線執行。")

    heading(doc, "八、科學誠實聲明")
    para(doc, "遊戲中「深層海水的 Mg²⁺ 經離子交換補入蛇紋石晶格空位（晶格修復），使 4–14 微米遠紅外線放射增強」是研究團隊的"
              "工作假設，尚待 FTIR、XRD、XPS 與放射率量測驗證；第 8 關的光譜為簡化示意圖，「修復後」曲線為預期結果，並非真實量測數據。"
              "遊戲中的遠紅外線只描述為「看不見的光／熱輻射」，本遊戲不宣稱遠紅外線或此材料對人體有任何功效。"
              "建議教師把這一點當作教學重點：科學家先提出猜想，再用實驗證據檢驗。")

    heading(doc, "九、相關連結")
    bullet(doc, ONLINE, bold_prefix="線上遊戲：")
    bullet(doc, REPO, bold_prefix="原始碼：")
    bullet(doc, "見「README_使用說明.md」（三種玩法、SmartScreen、相機權限、常見問題）。", bold_prefix="操作說明：")
    doc.core_properties.title = "墨玉晶格修復任務 WebAR 教師指引：課程與評量"
    doc.core_properties.author = CREDIT
    doc.core_properties.comments = ""
    doc.save(out_docx)


def to_pdf(docx_path, pdf_path):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(docx_path), ReadOnly=True)
        pages = d.ComputeStatistics(2)  # wdStatisticPages
        d.SaveAs2(str(pdf_path), FileFormat=17)
        d.Close(False)
    finally:
        word.Quit()
    return pages


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs"
    out.mkdir(parents=True, exist_ok=True)
    C = load_content()
    docx_path = (out / f"{NAME}.docx").resolve()
    pdf_path = (out / f"{NAME}.pdf").resolve()
    build(C, docx_path)
    pages = to_pdf(docx_path, pdf_path)
    print(f"docx: {docx_path}\npdf: {pdf_path}\npages: {pages}")


if __name__ == "__main__":
    main()
